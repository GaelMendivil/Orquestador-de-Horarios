"""Lee y escribe la configuración como libro de Excel (una hoja por tabla)."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .configuracion import (COLUMNAS, EJEMPLOS_FRANJAS, NOMBRES_TABLAS, _clave,
                            configuracion_vacia, normalizar)
from .entrada import ErrorInstancia

FUENTE = "Arial"
ENCABEZADO = PatternFill("solid", fgColor="D9E1F2")

AYUDA = {
    ("bloques", "Inicio"): "Hora de inicio del bloque, p. ej. 16:00.",
    ("bloques", "Fin"): "Hora de fin, p. ej. 17:40 (duración real de la clase).",
    ("bloques", "Turno"): "Turno al que pertenece el bloque (matutino, vespertino...). Los grupos usan este nombre.",
    ("salones", "Salón"): "Nombre o número del salón. No se repite.",
    ("salones", "Edificio"): "Opcional, sólo informativo por ahora.",
    ("salones", "Capacidad"): "Máximo de alumnos que caben. Una clase sólo va a salones donde cabe.",
    ("profesores", "Profesor"): "Nombre. Debe escribirse igual que en Materias.",
    ("profesores", "Disponible"): "Cuándo PUEDE dar clase (el algoritmo nunca lo rompe). Vacío = siempre. "
                                  + EJEMPLOS_FRANJAS,
    ("profesores", "Prefiere"): "Cuándo PREFIERE dar clase (se intenta respetar). Vacío = sin preferencia.",
    ("grupos", "Carrera"): "Clave o nombre de la carrera, p. ej. ICC.",
    ("grupos", "Semestre"): "Número de semestre (la generación).",
    ("grupos", "Grupo"): "Nombre del grupo dentro del semestre (A, B, G1...). Si sólo hay uno puede quedar vacío.",
    ("grupos", "Alumnos"): "Alumnos del grupo. Sus clases necesitan un salón con esa capacidad.",
    ("grupos", "Turno"): "Turno del grupo (de la hoja Bloques). Vacío = cualquier bloque del día.",
    ("materias", "Carrera"): "Carrera del plan de estudios.",
    ("materias", "Semestre"): "Semestre. La materia se abre para CADA grupo de esa carrera y semestre.",
    ("materias", "Materia"): "Nombre de la materia.",
    ("materias", "Profesor"): "Quién la imparte. Si hay varios, sepáralos con coma y se reparten entre los grupos.",
    ("materias", "Sesiones"): "Clases por semana (p. ej. 2 sesiones de 1:40).",
    ("materias", "Alumnos"): "Opcional: alumnos de la materia si no son todos los del grupo.",
}
PARA_QUE = {
    "General": "Nombre de la configuración y días de clase.",
    "Bloques": "Cómo se divide el día: un renglón por bloque de clase, con su turno.",
    "Salones": "Salones disponibles y su capacidad.",
    "Profesores": "Cuándo puede y cuándo prefiere dar clase cada profesor.",
    "Grupos": "Carreras y generaciones: cuántos grupos hay en cada semestre, sus alumnos y su turno.",
    "Materias": "Plan de estudios: qué materias lleva cada semestre de cada carrera y quién las da.",
}
FILAS_GENERAL = (
    ("nombre", "Nombre de la configuración", "Texto libre."),
    ("dias", "Días de clase por semana", "5 = lunes a viernes; 6 = lunes a sábado."),
)


def crear_plantilla(config: dict | None, ruta) -> None:
    """`ruta` puede ser un archivo abierto."""
    config = normalizar(config or configuracion_vacia())
    wb = Workbook()
    _hoja_instrucciones(wb.active)

    ws = wb.create_sheet("General")
    _encabezado(ws, ("Parámetro", "Valor", "Explicación"), (42, 40, 80))
    for i, (clave, etiqueta, ayuda) in enumerate(FILAS_GENERAL, 2):
        valor = config["general"].get(clave)
        ws.cell(i, 1, etiqueta).font = Font(name=FUENTE, bold=True)
        ws.cell(i, 2, valor).font = Font(name=FUENTE)
        ws.cell(i, 3, ayuda).font = Font(name=FUENTE, italic=True, color="595959")

    turnos = sorted({f["Turno"] for f in config["bloques"] if f["Turno"]})
    anchos = {"Disponible": 40, "Prefiere": 30, "Materia": 40, "Profesor": 30}
    for tabla, columnas in COLUMNAS.items():
        ws = wb.create_sheet(NOMBRES_TABLAS[tabla])
        _encabezado(ws, columnas, [anchos.get(c, 14) for c in columnas])
        for j, col in enumerate(columnas, 1):
            ws.cell(1, j).comment = Comment(AYUDA[(tabla, col)], "Organizador de horarios", width=300, height=110)
        for i, fila in enumerate(config[tabla], 2):
            for j, col in enumerate(columnas, 1):
                valor = fila[col]
                celda = ws.cell(i, j, valor)
                celda.font = Font(name=FUENTE)
                if col in ("Inicio", "Fin"):
                    celda.number_format = "@"
        ultima = len(config[tabla]) + 500
        if "Turno" in columnas and turnos and tabla != "bloques":
            letra = "ABCDEFG"[columnas.index("Turno")]
            _lista(ws, f"{letra}2:{letra}{ultima}", turnos, estricta=False)
        ws.freeze_panes = "A2"
    wb.save(ruta)


def _hoja_instrucciones(ws) -> None:
    ws.title = "Cómo llenar"
    ws["A1"] = "Configuración inicial del organizador de horarios"
    ws["A1"].font = Font(name=FUENTE, bold=True, size=14)
    ws["A2"] = ("Cada hoja es una de las tablas de la app, con las mismas columnas. Un renglón por elemento; "
                "los renglones vacíos se ignoran. Pasa el cursor sobre un encabezado para ver qué significa. "
                "Al terminar, sube este archivo en la app (Punto de partida → Subir archivo).")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 48
    _encabezado(ws, ("Hoja", "Para qué sirve"), (16, 100), fila=4)
    for i, (hoja, texto) in enumerate(PARA_QUE.items(), 5):
        ws.cell(i, 1, hoja).font = Font(name=FUENTE, bold=True)
        ws.cell(i, 2, texto).font = Font(name=FUENTE)
    ws.cell(12, 1, "Regla clave").font = Font(name=FUENTE, bold=True)
    ws.cell(12, 2, "Una materia de Materias se abre para cada grupo de Grupos con la misma carrera y semestre. "
                   "Ejemplo: ICC 7 tiene grupos G1 y G2 y 6 materias → 12 clases a programar.").font = Font(name=FUENTE)
    ws.cell(12, 2).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[12].height = 32


def _encabezado(ws, titulos, anchos, fila=1) -> None:
    for j, (titulo, ancho) in enumerate(zip(titulos, anchos), 1):
        c = ws.cell(fila, j, titulo)
        c.font = Font(name=FUENTE, bold=True)
        c.fill = ENCABEZADO
        ws.column_dimensions[c.column_letter].width = ancho


def _lista(ws, rango, opciones, estricta=True) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
    if not estricta:
        dv.errorStyle = "warning"
    ws.add_data_validation(dv)
    dv.add(rango)


def leer_excel(ruta) -> dict:
    wb = load_workbook(ruta, data_only=True)
    errores: list[str] = []
    config = configuracion_vacia()

    if "General" in wb.sheetnames:
        etiquetas = {_clave(etiqueta): clave for clave, etiqueta, _ in FILAS_GENERAL}
        for fila in wb["General"].iter_rows(min_row=2, values_only=True):
            clave = etiquetas.get(_clave(fila[0] or ""))
            if clave is None:
                continue
            valor = fila[1] if len(fila) > 1 else None
            config["general"][clave] = valor
    else:
        errores.append("Falta la hoja General")

    for tabla, columnas in COLUMNAS.items():
        nombre = NOMBRES_TABLAS[tabla]
        if nombre not in wb.sheetnames:
            errores.append(f"Falta la hoja {nombre}")
            continue
        filas = list(wb[nombre].iter_rows(values_only=True))
        if not filas:
            continue
        posicion = {_clave(h): j for j, h in enumerate(filas[0]) if h}
        faltan = [c for c in columnas if _clave(c) not in posicion]
        if faltan:
            errores.append(f"Hoja {nombre}: faltan las columnas {', '.join(faltan)}")
            continue
        for fila in filas[1:]:
            registro = {}
            for c in columnas:
                j = posicion[_clave(c)]
                valor = fila[j] if j < len(fila) else None
                if isinstance(valor, (dt.time, dt.datetime)):
                    valor = valor.strftime("%H:%M")
                registro[c] = valor
            config[tabla].append(registro)

    if errores:
        raise ErrorInstancia(errores)
    return normalizar(config)
